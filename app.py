import os,re,json,subprocess,tempfile,datetime,unicodedata,io
from pathlib import Path
from zoneinfo import ZoneInfo
from urllib.parse import urljoin
import requests
from bs4 import BeautifulSoup
from flask import Flask,request,jsonify,render_template
from werkzeug.utils import secure_filename

app=Flask(__name__)
app.config['MAX_CONTENT_LENGTH']=5*1024*1024
BASE=Path(__file__).parent
LOCAL=BASE/'state.json'
SUPA_URL=os.getenv('SUPABASE_URL','').rstrip('/')
SUPA_KEY=os.getenv('SUPABASE_SERVICE_ROLE_KEY','')
ADMIN_PASSWORD=os.getenv('ADMIN_PASSWORD','')
AB='https://abslusovice.cz/'
SUD='https://www.sudslusovice.cz/obedove-menu.html'
DAYS=['pondělí','úterý','středa','čtvrtek','pátek']
DEFAULT={'favorites':['řízek','svíčková','guláš'],'prerovska':{},'updated':None}

def load_state():
    if SUPA_URL and SUPA_KEY:
        try:
            r=requests.get(f'{SUPA_URL}/rest/v1/app_state',params={'id':'eq.main','select':'data'},headers={'apikey':SUPA_KEY,'Authorization':'Bearer '+SUPA_KEY},timeout=12)
            r.raise_for_status();rows=r.json()
            if rows:return {**DEFAULT,**rows[0]['data']}
        except Exception as e:app.logger.warning('DB read failed: %s',e)
    if LOCAL.exists():
        try:return {**DEFAULT,**json.loads(LOCAL.read_text(encoding='utf-8'))}
        except Exception:pass
    demo=BASE/'demo-data.json'
    if demo.exists():return {**DEFAULT,**json.loads(demo.read_text(encoding='utf-8'))}
    return DEFAULT.copy()

def save_state(data):
    if SUPA_URL and SUPA_KEY:
        r=requests.post(f'{SUPA_URL}/rest/v1/app_state',json={'id':'main','data':data},headers={'apikey':SUPA_KEY,'Authorization':'Bearer '+SUPA_KEY,'Content-Type':'application/json','Prefer':'resolution=merge-duplicates'},timeout=15)
        r.raise_for_status();return
    LOCAL.write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')

def today():return datetime.datetime.now(ZoneInfo('Europe/Prague')).date()
def clean(s):return re.sub(r'\s+',' ',s).strip()
def norm(s):return ''.join(c for c in unicodedata.normalize('NFD',s.lower()) if unicodedata.category(c)!='Mn')
def dish(name,restaurant,price=None,day=None,kind='main',note=''):
    return {'name':clean(name),'restaurant':restaurant,'price':price,'day':day,'kind':kind,'note':note}

def fetch_ab():
    html=requests.get(AB,timeout=15,headers={'User-Agent':'Mozilla/5.0 (compatible; MenuAggregator/1.0)'});html.raise_for_status()
    soup=BeautifulSoup(html.text,'html.parser')
    h=next((x for x in soup.find_all(['h2','h3','h4']) if 'Denní menu' in x.get_text(' ',strip=True)),None)
    if not h:raise ValueError('Nadpis denního menu nebyl nalezen')
    date_el=h.find_next(['h2','h3','h4']); m=re.search(r'(\d{1,2})\.(\d{1,2})\.(\d{4})',date_el.get_text(' ',strip=True)) if date_el else None
    if not m:raise ValueError('Nelze zjistit datum menu AB')
    date=datetime.date(int(m[3]),int(m[2]),int(m[1])).isoformat()
    # Extract only text between date heading and next main section.
    pieces=[]
    for e in date_el.next_elements:
        if getattr(e,'name',None) in ('h2',) and e is not date_el:break
        if getattr(e,'name',None) in ('li','p'):
            if e.find_parent(['li','p']) is None:pieces.append(e.get_text(' ',strip=True))
    if not pieces:
        section=[]
        for e in date_el.next_elements:
            if getattr(e,'name',None)=='h2' and e is not date_el:break
            if isinstance(e,str):section.append(str(e))
        pieces=[clean(x) for x in '\n'.join(section).splitlines()]
    out=[];kind='main'
    for p in pieces:
        p=clean(p)
        if not p:continue
        if p.lower()=='polévka':kind='soup';continue
        if p.lower()=='hlavní jídlo':kind='main';continue
        p=re.sub(r'^\d+[.)]\s*','',p)
        if re.search(r'\(\s*[\d, ]+\)',p) or (kind=='main' and len(p)>15 and not p.lower().startswith('kde nás')):
            p=re.sub(r'\s*\*?\s*\([\d, ]+\)\s*$','',p).strip(' *')
            if p and not any(x['name']==p for x in out):out.append(dish(p,'AB Slušovice',day=date,kind=kind))
    if not out:raise ValueError('Z webu AB se nepodařilo získat položky')
    return out,date

def parse_pdf(data):
    try:
        import fitz
        doc=fitz.open(stream=data,filetype='pdf')
        return '\n'.join(p.get_text(sort=True) for p in doc)
    except Exception:return ''

def fetch_sud():
    r=requests.get(SUD,timeout=15);r.raise_for_status();s=BeautifulSoup(r.text,'html.parser')
    a=next((a for a in s.find_all('a',href=True) if '.pdf' in a['href'].lower()),None)
    if not a:raise ValueError('PDF jídelníček SUD nebyl nalezen')
    pdfurl=urljoin(SUD,a['href']);pr=requests.get(pdfurl,timeout=20);pr.raise_for_status()
    raw=parse_pdf(pr.content)
    if len(raw.strip())<40:raise ValueError('PDF SUD neobsahuje čitelný text (pravděpodobně obrázek)')
    return parse_generic_menu(raw,'SUD Slušovice'),pdfurl

def parse_generic_menu(text,restaurant):
    lines=[clean(x) for x in text.splitlines() if clean(x)]
    out=[];current=None;week_start=None
    for line in lines:
        w=re.search(r'\b(pondělí|úterý|středa|čtvrtek|pátek)\b\s*(\d{1,2})\s*[.]\s*(\d{1,2})\s*[.]?\s*(\d{4})?',line,re.I)
        if w:
            y=int(w[4]) if w[4] else today().year
            try:current=datetime.date(y,int(w[3]),int(w[2])).isoformat()
            except ValueError:current=None
            line=clean(line[w.end():])
        if not current:continue
        if not line:continue
        if re.match(r'^(polévka|polevka)\s*:',line,re.I):
            name=clean(line.split(':',1)[1]);
            if name:out.append(dish(re.sub(r'\s*\([\d, ]+\)$','',name),restaurant,day=current,kind='soup'))
            continue
        # Lines with numbered menu entries, or explicit CZK prices.
        match=re.match(r'^(?:\d+[.)]\s*)?(?:(?:\d+\s*g|\d+\s*ks)\s+)?(.+?)\s+(\d{2,3})\s*(?:Kč|,-)\s*$',line,re.I)
        numbered=re.match(r'^\d+[.)]\s*(?:(?:\d+\s*g|\d+\s*ks)\s+)?(.+)$',line,re.I)
        if match:out.append(dish(re.sub(r'\s*\([\d, ]+\)$','',match[1]),restaurant,int(match[2]),current))
        elif numbered:out.append(dish(re.sub(r'\s*\([\d, ]+\)$','',numbered[1]),restaurant,day=current))
    return out

def parse_prerovska(text):
    # Word extraction keeps paragraphs but often splits day headings from menu entries.
    lines=[clean(x) for x in text.splitlines() if clean(x)]
    out=[];current=None;kind='main';year=today().year
    ym=re.search(r'(20\d\d)',text)
    if ym:year=int(ym[1])
    for line in lines:
        d=re.search(r'\b(pondělí|úterý|středa|čtvrtek|pátek)\b\s*(\d{1,2})\s*[.]\s*(\d{1,2})',line,re.I)
        if d:
            try:current=datetime.date(year,int(d[3]),int(d[2])).isoformat()
            except ValueError:current=None
            line=clean(line[d.end():]);kind='main'
        if not current:continue
        if re.search(r'\bpolévka\s*:',line,re.I):
            name=clean(re.split(r'polévka\s*:',line,flags=re.I)[-1]);name=re.sub(r'\s*\([\d, ]+\).*','',name)
            if name:out.append(dish(name,'Přerovská restaurace',day=current,kind='soup'))
            continue
        match=re.match(r'^([123])[.)]\s*(?:(?:\d+\s*g|\d+\s*ks)\s+)?(.+)$',line,re.I)
        if match:
            name=re.sub(r'\s*\([\d, ]+\).*','',match[2]);out.append(dish(name,'Přerovská restaurace',165,current))
        elif out and out[-1]['day']==current and out[-1]['kind']=='main' and line.startswith('('):
            pass
    return out

def word_text(data,filename):
    ext=Path(filename).suffix.lower()
    with tempfile.TemporaryDirectory() as tmp:
        path=Path(tmp)/('menu'+ext);path.write_bytes(data)
        if ext=='.doc':
            result=subprocess.run(['antiword','-m','UTF-8.txt','-w','0',str(path)],capture_output=True,text=True,timeout=20)
            if result.returncode:raise ValueError('Soubor DOC nelze přečíst: '+result.stderr[:100])
            return result.stdout
        if ext=='.docx':
            from docx import Document
            doc=Document(io.BytesIO(data));return '\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(' '.join(c.text for c in row.cells) for t in doc.tables for row in t.rows)
        raise ValueError('Povolené soubory jsou .doc a .docx')

def authorized():
    supplied=request.headers.get('X-Admin-Password','')
    import hmac
    return bool(ADMIN_PASSWORD) and hmac.compare_digest(supplied,ADMIN_PASSWORD)

@app.get('/')
def home():return render_template('index.html')

@app.get('/api/menu')
def menu():
    date=request.args.get('date') or today().isoformat()
    try:datetime.date.fromisoformat(date)
    except ValueError:return jsonify(error='Neplatné datum'),400
    state=load_state();items=[x for x in state.get('prerovska',{}).get('items',[]) if x['day']==date]
    sources={'prerovska':{'status':'ok' if state.get('prerovska',{}).get('items') else 'missing','updated':state.get('updated')}}
    try:
        ab,abdate=fetch_ab()
        if abdate==date:items+=ab;sources['ab']={'status':'ok','date':abdate}
        else:sources['ab']={'status':'other_date','date':abdate}
    except Exception as e:sources['ab']={'status':'error','message':str(e)}
    try:
        sud,pdfurl=fetch_sud()
        items += [x for x in sud if x['day']==date]
        sources['sud']={'status':'ok' if sud else 'unreadable','url':pdfurl}
    except Exception as e:sources['sud']={'status':'error','message':str(e)}
    fav=state.get('favorites',[])
    for x in items:x['favorite']=any(norm(f) in norm(x['name']) for f in fav if f.strip())
    items.sort(key=lambda x:(not x['favorite'],x['kind']!='main',x['restaurant'],x['name']))
    return jsonify(date=date,items=items,favorites=fav,sources=sources)

@app.post('/api/admin/upload')
def upload():
    if not authorized():return jsonify(error='Nesprávné heslo správce'),403
    f=request.files.get('file')
    if not f or Path(f.filename).suffix.lower() not in ('.doc','.docx'):return jsonify(error='Nahraj DOC nebo DOCX'),400
    try:
        text=word_text(f.read(),f.filename);items=parse_prerovska(text)
        if len(items)<3:raise ValueError('Z dokumentu se nepodařilo rozpoznat dost jídel; zkontroluj jeho strukturu.')
        state=load_state();state['prerovska']={'items':items,'filename':secure_filename(f.filename)};state['updated']=datetime.datetime.now(ZoneInfo('Europe/Prague')).isoformat();save_state(state)
        return jsonify(ok=True,count=len(items),days=sorted(set(x['day'] for x in items)))
    except Exception as e:return jsonify(error=str(e)),400

@app.post('/api/admin/favorites')
def favorites():
    if not authorized():return jsonify(error='Nesprávné heslo správce'),403
    terms=request.json.get('favorites',[])
    if not isinstance(terms,list) or len(terms)>50 or any(not isinstance(t,str) or len(t)>60 for t in terms):return jsonify(error='Neplatné preference'),400
    state=load_state();state['favorites']=[t.strip() for t in terms if t.strip()];save_state(state)
    return jsonify(ok=True,favorites=state['favorites'])

if __name__=='__main__':app.run(debug=True,port=5000)
